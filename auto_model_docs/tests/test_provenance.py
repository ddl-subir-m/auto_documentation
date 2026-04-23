"""Tests for provenance stamping (8 fields: 5 MVP + 3 phase-2)."""

from __future__ import annotations

import json
import os
import re
import sqlite3
import subprocess
import sys
from pathlib import Path
from unittest.mock import patch

import pytest
from docx import Document

from autodoc import provenance


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

def _read_row(db_path: str, provenance_id: str) -> dict | None:
    with sqlite3.connect(db_path) as conn:
        conn.row_factory = sqlite3.Row
        cur = conn.execute(
            "SELECT * FROM provenance WHERE provenance_id = ?", (provenance_id,)
        )
        row = cur.fetchone()
    return dict(row) if row else None


def _make_record(**overrides) -> provenance.ProvenanceRecord:
    defaults = dict(
        provenance_id="abc123def456",
        bundle_id="bundle-xyz",
        policy_version_id="policy-v42",
        commit_sha="deadbeef" * 5,
        generated_by_user="alice",
        generated_at="2026-04-22T12:34:56Z",
        generator_version="0.1.0",
        template_version="1.0",
        run_environment={
            "domino_env_id": "env-123",
            "domino_run_id": "run-456",
            "python_version": "3.11.14",
        },
    )
    defaults.update(overrides)
    return provenance.ProvenanceRecord(**defaults)


def _fs_commit(temp_path: str, final_path: str) -> None:
    os.makedirs(os.path.dirname(final_path) or ".", exist_ok=True)
    os.replace(temp_path, final_path)


# ---------------------------------------------------------------------------
# Custom-property stamping
# ---------------------------------------------------------------------------

def test_all_8_fields_appear_in_word_custom_properties(tmp_path):
    db_path = str(tmp_path / "prov.db")
    final_path = str(tmp_path / "out.docx")
    record = _make_record()

    doc = Document()
    doc.add_paragraph("Hello")
    provenance.save_with_provenance(doc, final_path, record, db_path, _fs_commit)

    props = provenance.read_custom_properties(final_path)
    assert props["bundle_id"] == "bundle-xyz"
    assert props["policy_version_id"] == "policy-v42"
    assert props["commit_sha"] == "deadbeef" * 5
    assert props["generated_by_user"] == "alice"
    assert props["generated_at"] == "2026-04-22T12:34:56Z"
    assert props["generator_version"] == "0.1.0"
    assert props["template_version"] == "1.0"
    # run_environment is serialized as JSON
    run_env = json.loads(props["run_environment"])
    assert run_env["domino_env_id"] == "env-123"
    assert run_env["domino_run_id"] == "run-456"
    assert run_env["python_version"] == "3.11.14"

    assert set(props.keys()) == set(provenance.PROVENANCE_FIELDS)


def test_all_8_fields_present_in_sqlite_row(tmp_path):
    db_path = str(tmp_path / "prov.db")
    final_path = str(tmp_path / "out.docx")
    record = _make_record()

    doc = Document()
    provenance.save_with_provenance(doc, final_path, record, db_path, _fs_commit)

    row = _read_row(db_path, record.provenance_id)
    assert row is not None
    assert row["bundle_id"] == "bundle-xyz"
    assert row["policy_version_id"] == "policy-v42"
    assert row["commit_sha"] == "deadbeef" * 5
    assert row["generated_by_user"] == "alice"
    assert row["generated_at"] == "2026-04-22T12:34:56Z"
    assert row["generator_version"] == "0.1.0"
    assert row["template_version"] == "1.0"
    assert row["docx_path"] == final_path
    assert isinstance(row["created_ts"], float)

    run_env = json.loads(row["run_environment"])
    assert run_env == {
        "domino_env_id": "env-123",
        "domino_run_id": "run-456",
        "python_version": "3.11.14",
    }


# ---------------------------------------------------------------------------
# capture_context: commit SHA, user, phase-2 fields
# ---------------------------------------------------------------------------

def test_provenance_id_unique_across_runs(monkeypatch):
    monkeypatch.setenv("DOMINO_STARTING_USERNAME", "bob")
    ids = {provenance.capture_context().provenance_id for _ in range(50)}
    assert len(ids) == 50
    for pid in ids:
        assert re.fullmatch(r"[0-9a-f]{12}", pid)


def test_commit_sha_fallback_when_git_fails(monkeypatch):
    def _raise(*_a, **_kw):
        raise FileNotFoundError("git not installed")

    monkeypatch.setenv("DOMINO_WORKING_DIR_GIT_REF_ID", "env-fallback-sha")
    monkeypatch.delenv("DOMINO_STARTING_USERNAME", raising=False)
    with patch.object(subprocess, "run", side_effect=_raise):
        record = provenance.capture_context()

    assert record.commit_sha == "env-fallback-sha"
    assert record.generated_by_user == "unknown"


def test_commit_sha_none_when_git_fails_and_no_env(monkeypatch):
    def _raise(*_a, **_kw):
        raise subprocess.CalledProcessError(128, "git")

    monkeypatch.delenv("DOMINO_WORKING_DIR_GIT_REF_ID", raising=False)
    with patch.object(subprocess, "run", side_effect=_raise):
        record = provenance.capture_context()

    assert record.commit_sha is None


def test_generated_by_user_fallback_to_unknown(monkeypatch):
    monkeypatch.delenv("DOMINO_STARTING_USERNAME", raising=False)
    record = provenance.capture_context()
    assert record.generated_by_user == "unknown"


def test_generator_version_read_from_pyproject():
    """generator_version matches [project].version in the package's pyproject.toml."""
    record = provenance.capture_context()
    assert record.generator_version and record.generator_version != "unknown"
    assert re.fullmatch(r"\d+\.\d+\.\d+", record.generator_version)


def test_generator_version_fallback_when_pyproject_missing(monkeypatch):
    fake_root = Path("/nonexistent/path/does/not/exist")
    with patch.object(provenance, "__file__", str(fake_root / "provenance.py")):
        version = provenance._resolve_generator_version()
    assert version == "unknown"


def test_run_environment_captured(monkeypatch):
    monkeypatch.setenv("DOMINO_ENVIRONMENT_ID", "env-abc")
    monkeypatch.setenv("DOMINO_RUN_ID", "run-xyz")
    record = provenance.capture_context()
    assert record.run_environment["domino_env_id"] == "env-abc"
    assert record.run_environment["domino_run_id"] == "run-xyz"
    assert record.run_environment["python_version"] == (
        f"{sys.version_info.major}.{sys.version_info.minor}.{sys.version_info.micro}"
    )


def test_run_environment_env_missing_is_none(monkeypatch):
    monkeypatch.delenv("DOMINO_ENVIRONMENT_ID", raising=False)
    monkeypatch.delenv("DOMINO_RUN_ID", raising=False)
    record = provenance.capture_context()
    assert record.run_environment["domino_env_id"] is None
    assert record.run_environment["domino_run_id"] is None


# ---------------------------------------------------------------------------
# template_version: comes from the spec; required
# ---------------------------------------------------------------------------

def test_template_version_extracted_from_spec_dict():
    record = provenance.capture_context(spec={"template_version": "1.0", "title": "x"})
    assert record.template_version == "1.0"


def test_template_version_extracted_from_object_attr():
    class FakeSpec:
        template_version = "2.5"

    record = provenance.capture_context(spec=FakeSpec())
    assert record.template_version == "2.5"


def test_missing_template_version_in_spec_raises_clean_error():
    with pytest.raises(ValueError, match="template_version"):
        provenance.capture_context(spec={"title": "no template version here"})


def test_empty_template_version_in_spec_raises():
    with pytest.raises(ValueError, match="template_version"):
        provenance.capture_context(spec={"template_version": "   "})


def test_template_version_defaults_to_unknown_when_no_spec():
    """Backwards-compat: callers that don't supply a spec get 'unknown'."""
    record = provenance.capture_context()
    assert record.template_version == "unknown"


# ---------------------------------------------------------------------------
# Atomicity
# ---------------------------------------------------------------------------

def test_atomicity_rolls_back_row_on_commit_failure(tmp_path):
    db_path = str(tmp_path / "prov.db")
    final_path = str(tmp_path / "never_written.docx")
    record = _make_record(provenance_id="rollback123")

    def _failing_commit(_temp, _final):
        raise RuntimeError("simulated write failure")

    doc = Document()
    with pytest.raises(RuntimeError, match="simulated write failure"):
        provenance.save_with_provenance(doc, final_path, record, db_path, _failing_commit)

    assert _read_row(db_path, "rollback123") is None
    assert not Path(final_path).exists()


def test_capture_context_with_none_bundle_and_policy(monkeypatch):
    monkeypatch.setenv("DOMINO_STARTING_USERNAME", "carol")
    record = provenance.capture_context(bundle_id=None, policy_version_id=None)
    assert record.bundle_id is None
    assert record.policy_version_id is None
    assert record.generated_by_user == "carol"
    assert re.fullmatch(r"\d{4}-\d{2}-\d{2}T\d{2}:\d{2}:\d{2}Z", record.generated_at)


def test_none_fields_stamp_as_empty_strings(tmp_path):
    db_path = str(tmp_path / "prov.db")
    final_path = str(tmp_path / "out.docx")
    record = _make_record(bundle_id=None, policy_version_id=None, commit_sha=None)

    doc = Document()
    provenance.save_with_provenance(doc, final_path, record, db_path, _fs_commit)

    props = provenance.read_custom_properties(final_path)
    assert props["bundle_id"] == "" or props["bundle_id"] is None
    assert props["policy_version_id"] == "" or props["policy_version_id"] is None

    row = _read_row(db_path, record.provenance_id)
    assert row["bundle_id"] is None
    assert row["policy_version_id"] is None
    assert row["commit_sha"] is None


def test_default_db_path_respects_env(monkeypatch, tmp_path):
    override = str(tmp_path / "custom.db")
    monkeypatch.setenv("AUTODOC_PROVENANCE_DB", override)
    assert provenance.default_db_path() == override


# ---------------------------------------------------------------------------
# Schema migration: pre-phase-2 DB (5 MVP cols only) upgrades cleanly
# ---------------------------------------------------------------------------

_PRE_PHASE2_SCHEMA = """
CREATE TABLE provenance (
    provenance_id TEXT PRIMARY KEY,
    bundle_id TEXT,
    policy_version_id TEXT,
    commit_sha TEXT,
    generated_by_user TEXT NOT NULL,
    generated_at TEXT NOT NULL,
    docx_path TEXT NOT NULL,
    created_ts REAL NOT NULL
);
"""


def test_migration_upgrades_pre_phase2_schema(tmp_path):
    """A DB created before the 3 phase-2 columns existed is upgraded in place."""
    db_path = str(tmp_path / "legacy.db")

    with sqlite3.connect(db_path) as conn:
        conn.executescript(_PRE_PHASE2_SCHEMA)
        conn.execute(
            """
            INSERT INTO provenance
                (provenance_id, bundle_id, policy_version_id, commit_sha,
                 generated_by_user, generated_at, docx_path, created_ts)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?)
            """,
            ("legacy001", "b1", "p1", "sha1", "alice", "2026-01-01T00:00:00Z",
             "/x.docx", 1700000000.0),
        )
        conn.commit()

    # Re-open via the library; migration should add the 3 new columns without loss.
    record = _make_record(provenance_id="new001")
    provenance.write_provenance_row(db_path, record, "/new.docx")

    with sqlite3.connect(db_path) as conn:
        cols = {row[1] for row in conn.execute("PRAGMA table_info(provenance)")}

    assert {"generator_version", "template_version", "run_environment"} <= cols

    legacy = _read_row(db_path, "legacy001")
    assert legacy is not None
    assert legacy["bundle_id"] == "b1"
    # New columns default to NULL for legacy rows.
    assert legacy["generator_version"] is None
    assert legacy["template_version"] is None
    assert legacy["run_environment"] is None

    new_row = _read_row(db_path, "new001")
    assert new_row["generator_version"] == "0.1.0"
    assert new_row["template_version"] == "1.0"
    assert json.loads(new_row["run_environment"])["python_version"] == "3.11.14"


def test_migration_is_idempotent(tmp_path):
    """Running the migration repeatedly must not error or duplicate columns."""
    db_path = str(tmp_path / "repeat.db")
    for _ in range(3):
        with provenance._connect(db_path) as conn:
            cols = [row[1] for row in conn.execute("PRAGMA table_info(provenance)")]
        assert cols.count("generator_version") == 1
        assert cols.count("template_version") == 1
        assert cols.count("run_environment") == 1
